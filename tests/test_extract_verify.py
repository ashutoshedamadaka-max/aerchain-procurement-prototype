"""Offline tests for the deterministic verifier (no model calls)."""
import json

import openpyxl
import pytest

from src.extract.verify import Rfx, verify_row, verify_terms
from src.extract.workbook import display


@pytest.fixture
def env(tmp_path):
    (tmp_path / "rfx.json").write_text(json.dumps({"rfx_id": "R", "evaluation_date": "2026-03-16", "invited_vendors": [], "lines": [
        {"line_no": 1, "length_mm": 300, "width_mm": 200, "height_mm": 150, "ply": 3, "liner_gsm": 150, "print_spec": "plain"}]}))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q"
    ws.append(["300x200x150", 3, "150/120/150", "plain", 1000, "10.00", 10000])
    ws["H1"], ws["B5"], ws["C5"] = "x", 25000000, "Add 3%"
    wb.save(tmp_path / "v.xlsx")
    from src.extract.workbook import Workbook
    return Workbook(tmp_path / "v.xlsx"), Rfx(tmp_path / "rfx.json")


def cell(a, s):
    return {"anchor": a, "snippet": s}


def row(**over):
    r = {"row": 1, "sheet": "Q", "size": cell("Q!A1", "300x200x150"), "ply": cell("Q!B1", "3"), "gsm_stack": cell("Q!C1", "150/120/150"),
         "printing": cell("Q!D1", "plain"), "annual_qty": cell("Q!E1", "1000"), "rate_status": "priced", "rate": cell("Q!F1", "10.00"),
         "amount": cell("Q!G1", "10000"), "uom_text": "Nos", "basis": "per_piece", "currency": "INR"}
    return {**r, **over}


def codes(c):
    return [f.code for f in c.failures]


def test_clean_row_passes(env):
    assert codes(verify_row(row(), *env)) == []


def test_arithmetic_mismatch_flags_both_price_and_total(env):
    env[0].wb["Q"]["G1"] = 1000
    c = verify_row(row(amount=cell("Q!G1", "1000")), *env)
    assert codes(c) == ["arith_mismatch"] and set(c.failures[0].fields) == {"unit_price", "line_total"}


def test_out_of_band_rate(env):
    env[0].wb["Q"]["F1"], env[0].wb["Q"]["G1"] = "1000.00", 1000000
    assert "price_out_of_band" in codes(verify_row(row(rate=cell("Q!F1", "1000.00"), amount=cell("Q!G1", "1,000,000")), *env))


def test_snippet_not_verbatim_fails_provenance(env):
    assert "provenance_unverified" in codes(verify_row(row(rate=cell("Q!F1", "10.50")), *env))


def test_wrong_row_anchor_fails_provenance(env):
    assert "provenance_unverified" in codes(verify_row(row(rate=cell("Q!F2", "10.00")), *env))


def test_gsm_deviation_is_a_finding_not_a_failure(env):
    env[0].wb["Q"]["C1"] = "120/120/120"
    c = verify_row(row(gsm_stack=cell("Q!C1", "120/120/120")), *env)
    assert codes(c) == [] and c.deviations["unit_price"] == "spec_mismatch_gsm"


def test_unmatched_size_fails_matching(env):
    env[0].wb["Q"]["A1"] = "999x999x999"
    assert "line_unmatched" in codes(verify_row(row(size=cell("Q!A1", "999x999x999")), *env))


def test_declined_row_is_not_a_failure(env):
    assert codes(verify_row(row(rate_status="declined", rate=None, amount=None), *env)) == []


def test_large_general_numbers_display_without_exponent(env):
    assert display(env[0].wb["Q"]["B5"]) == "25000000"


def test_tier_lower_bound_must_parse(env):
    wb = env[0]
    tier = {"min_value": cell("Q!B5", "2.5e+07"), "max_value": None, "effect": cell("Q!C5", "Add 3%"), "effect_kind": "uplift_pct"}
    data = {"vendor_name": None, "quotation_date": None, "incoterm": None, "conditions": [], "tiers": [tier]}
    assert any("lower bound" in p or "provenance" in p for p in verify_terms(data, wb))


# ---- multi-format additions ----
import pathlib

from src.extract.sources import DocxSource, EmailSource, PdfSource
from src.extract.verify import parse_amount_inr, parse_int, single_int

ART = pathlib.Path(__file__).resolve().parents[1] / "dataset" / "artifacts"


def test_number_words_and_ambiguity():
    assert parse_int("This offer is valid for ten days.") == 10
    assert single_int("Payment: 30 days net.") == 30
    assert single_int("A further 2% discount above Rs. 50 lakh. Payment: 30 days net.") is None


def test_indian_amounts_are_computed_in_code():
    assert parse_amount_inr("Rs. 50 lakh") == 5_000_000 and parse_amount_inr("2.5 crore") == 25_000_000 and parse_amount_inr("Rs 7,50,000") == 750000


def test_missing_gsm_is_spec_incomplete(env):
    assert "spec_incomplete" in codes(verify_row(row(gsm_stack=None), *env))


def test_verbatim_basis_must_agree_with_the_models_reading(env):
    r = row(basis_text=cell("Q!D1", "plain"), basis="per_100_pieces")
    env[0].wb["Q"]["D1"] = "per 100 pcs"
    assert "basis_unclear" in codes(verify_row(row(basis_text=cell("Q!D1", "per 100 pcs"), basis="per_piece"), *env))


def test_unclear_currency_escalates(env):
    assert "currency_unclear" in codes(verify_row(row(currency="unclear"), *env))


@pytest.mark.skipif(not (ART / "V2_response.pdf").exists(), reason="artifact not generated")
def test_pdf_wrapped_basis_is_rejoined_and_locators_canonicalise():
    p = PdfSource(ART / "V2_response.pdf")
    assert len(p.candidate_rows()) == 30 and p.text_at("p1:r10:desc").endswith("per 100 pcs")
    assert p.canonical("r20:desc", "p2") == "p2:r20:desc" and p.canonical("doc!p2:L4", "p2") == "p2:L4"
    assert any(size == 7.0 and "settlement discount" in t for t, size in p.free.values())


@pytest.mark.skipif(not (ART / "V3_response.docx").exists(), reason="artifact not generated")
def test_docx_paragraph_locators_and_email_lines():
    d = DocxSource(ART / "V3_response.docx")
    assert d.row_of("para29") == ("doc", 29) and d.canonical("doc!para15") == "para15" and len(d.candidate_rows()) == 28
    e = EmailSource(ART / "V5_response.txt")
    assert e.verbatim("L6", "rest same as last year") and not e.verbatim("L6", "Rs 38")
