#!/usr/bin/env python3
"""Render deliberately inconsistent vendor documents from frozen SQLite truth.

Run: python scripts/generate_vendor_docs.py [--seed 20260311]
Requires: openpyxl, reportlab, python-docx. Reads truth in SQLite read-only mode.
Only the four named responses in dataset/artifacts are written; V4 is untouched.
Seed controls cosmetic SKU letters/UOM choices, never quoted values.

Truth takes precedence over the superseded brief: V1 has four award-value slabs;
V2 has 30-day net terms and no 15-day early-payment requirement. Preserve the
stored line totals even where wrong. V3 para29 deliberately mentions only the
first declined size; the other two are absent, per the requested format trap.
Planned V1 cells, V2 quote pages and V3 commercial paragraph numbers are retained.
Requested filenames differ from submissions.file_name; truth is never updated.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import random
import re
import sqlite3
import string
import sys
from xml.sax.saxutils import escape
import zipfile

ROOT = Path(__file__).resolve().parents[1]
SEED = 20260311
sys.path.insert(0, str(ROOT / "scripts"))
from dataset_text import ATTACHMENTS  # noqa: E402

ATTACH_FILE = {a[0]: a[3] for a in ATTACHMENTS}     # attachment id -> the certificate file the vendor sent


def read_truth(root: Path) -> dict:
    path = root / "dataset/truth/truth.sqlite"
    conn = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        tables = {
            table: [dict(r) for r in conn.execute(f"SELECT * FROM {table}")]
            for table in ("rfx_lines", "vendors", "submissions", "bid_fields",
                          "tier_rules", "conditions", "questionnaire_answers")
        }
    finally:
        conn.close()
    tables["outcomes"] = json.loads((root / "dataset/truth/outcomes.json").read_text())
    return tables


def vendor_rows(data: dict, table: str, vendor: str) -> list[dict]:
    sub = next(s["submission_id"] for s in data["submissions"] if s["vendor_id"] == vendor)
    return [r for r in data[table] if r.get("vendor_id") == vendor or r.get("submission_id") == sub]


def field(data: dict, vendor: str, line: int, name: str) -> dict:
    return next(r for r in vendor_rows(data, "bid_fields", vendor)
                if r["rfx_line_no"] == line and r["field_name"] == name)


def company(data: dict, vendor: str) -> str:
    return next(v["name"] for v in data["vendors"] if v["vendor_id"] == vendor)


def submission(data: dict, vendor: str) -> dict:
    return next(s for s in data["submissions"] if s["vendor_id"] == vendor)


def dims(line: dict) -> str:
    return "x".join(str(line[k]) for k in ("length_mm", "width_mm", "height_mm"))


def stable_zip(path: Path, stamp: datetime) -> None:
    """Remove packaging timestamps so identical truth and seed give identical bytes."""
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as dest:
        for info, body in entries:
            if info.filename == "docProps/core.xml":
                body = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    lambda m: m[1] + (stamp.isoformat() + "Z").encode() + m[2],
                    body,
                )
            info.date_time = stamp.timetuple()[:6]
            dest.writestr(info, body)


def make_excel(data: dict, out: Path, rng: random.Random) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    stamp = datetime.fromisoformat(submission(data, "V1")["received_on"])
    wb.properties.created = wb.properties.modified = stamp
    wb.properties.creator = company(data, "V1")
    ws.merge_cells("A1:K2")
    ws["A1"] = company(data, "V1")
    ws["A1"].font = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="244B42")
    ws.merge_cells("A3:K3")
    ws["A3"] = "COMMERCIAL OFFER / CORRUGATED SHIPPERS"
    ws.merge_cells("A4:F4")
    ws["A4"] = "To: Purchase Department"
    ws.merge_cells("G4:K4")
    ws["G4"] = submission(data, "V1")["received_on"]
    ws.merge_cells("A5:K5")
    ws["A5"] = "Ex-works | Rates and board construction as below"
    # Header at row 7 and two genuinely empty rows 8-9 before the data.
    headers = ["Our article", "Box size mm", "Ply", "Flute", "GSM stack", "Printing",
               "Annual qty", "Rate INR", "UOM", "BF", "Amount INR"]
    for col, title in enumerate(headers, 1):
        ws.cell(7, col, title)
    codes = set()
    uoms = ["Nos", "nos.", "PCS"]
    rng.shuffle(uoms)
    for ln in sorted(data["rfx_lines"], key=lambda x: x["line_no"]):
        n = ln["line_no"]
        p = field(data, "V1", n, "unit_price")
        total = field(data, "V1", n, "line_total")
        spec = field(data, "V1", n, "declared_liner_gsm")
        row = int(re.search(r"\d+$", p["anchor"]).group())
        while True:
            code = "SP-" + "".join(rng.choices(string.ascii_uppercase, k=5))
            if code not in codes:
                codes.add(code)
                break
        # The entire rate column is text; no correction or recomputation of line totals.
        values = [code, dims(ln), ln["ply"], ln["flute"], spec["snippet"],
                  ln["print_spec"], ln["annual_qty"], p["snippet"],
                  uoms[(n - 1) % len(uoms)],
                  ln["bf"], total["value"]]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 8).number_format = "@"
        ws.cell(row, 11).number_format = '#,##0.00'
        ws.cell(row, 7).number_format = '#,##0'
        ws.row_dimensions[row].height = 25
    # A formula total, inside the filter/data range, sums the stated (including bad) totals.
    last = max(int(re.search(r"\d+$", r["anchor"]).group())
               for r in vendor_rows(data, "bid_fields", "V1") if r["field_name"] == "line_total")
    ws.cell(last + 1, 2, "TOTAL")
    ws.cell(last + 1, 11, f"=SUM(K10:K{last})")
    ws.cell(last + 1, 11).number_format = '#,##0.00'
    ws.auto_filter.ref = f"A7:K{last + 1}"
    ws.freeze_panes = "C10"
    ws.print_title_rows = "1:9"
    widths = [17, 23, 7, 8, 31, 19, 15, 13, 9, 7, 20]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True

    terms = wb.create_sheet("Terms & Conditions")
    terms.merge_cells("B1:D1")
    terms["B1"] = "Commercial terms"
    for r in vendor_rows(data, "conditions", "V1"):
        cell = r["anchor"].split("!")[1]
        terms[cell] = r["snippet"]
        terms.merge_cells(start_row=terms[cell].row, start_column=2, end_row=terms[cell].row, end_column=4)
        terms[cell].alignment = Alignment(wrap_text=True, vertical="center")
        terms.row_dimensions[terms[cell].row].height = 32
    for c, title in zip("BCD", ["Awarded value from INR", "Below INR", "Rate adjustment"]):
        terms[f"{c}9"] = title
    for row, tier in enumerate(vendor_rows(data, "tier_rules", "V1"), 10):
        terms.cell(row, 2, tier["min_value_inr"])
        terms.cell(row, 3, tier["max_value_inr"] if tier["max_value_inr"] is not None else "and above")
        terms.cell(row, 4, "Rates as quoted" if tier["effect_pct"] == 0 else f"Add {tier['effect_pct']:g}%")
    terms.merge_cells("B15:D16")
    terms["B15"] = "Slabs apply to the value actually awarded; lower bound inclusive, upper bound exclusive."
    terms["B15"].alignment = Alignment(wrap_text=True)
    for col, width in {"A": 3, "B": 28, "C": 58, "D": 64}.items():
        terms.column_dimensions[col].width = width

    # The supplier questionnaire has its own sheet: V1 is the organised converter, and answers it where the buyer's form asked for it.
    quest = wb.create_sheet("Questionnaire")
    quest.append(["Ref", "Question", "Our response", "Attachment"])
    for row, answer in enumerate(vendor_rows(data, "questionnaire_answers", "V1"), 2):
        quest.append([answer["q_no"], answer["question"], answer["answer_text"], ATTACH_FILE.get(answer["attachment_id"])])
        quest.row_dimensions[row].height = 48
    for col, width in zip("ABCD", [6, 62, 58, 26]):
        quest.column_dimensions[col].width = width

    board = wb.create_sheet("Board Specification")
    board.append(["Our article", "Size mm", "Ply", "Board GSM stack", "Flute", "BF"])
    for row in range(10, last + 1):
        board.append([ws.cell(row, c).value for c in (1, 2, 3, 5, 4, 10)])
    for col, width in zip("ABCDEF", [19, 25, 10, 36, 12, 10]):
        board.column_dimensions[col].width = width
    wb.move_sheet(quest, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Questionnaire"))     # the fourth sheet
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                if cell.value is not None and cell.coordinate != "A1":
                    cell.font = Font(name="Calibri", size=11, color="243C34")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
        header = 7 if sheet == ws else 9 if sheet == terms else 1
        for cell in sheet[header]:
            cell.fill = PatternFill("solid", fgColor="D9E8DF")
            cell.font = Font(name="Calibri", bold=True, size=11)
            cell.border = Border(bottom=Side(style="thin", color="537867"))
    path = out / "V1_response.xlsx"
    wb.save(path)
    stable_zip(path, stamp)


def make_pdf(data: dict, out: Path) -> None:
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle

    canvas = Canvas(str(out / "V2_response.pdf"), pagesize=A4, invariant=1)
    canvas.setTitle("Commercial quotation")
    canvas.setAuthor(company(data, "V2"))
    width, height = A4
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8, leading=10)
    tiny = ParagraphStyle("foot", parent=body, fontSize=7, leading=9)
    question = ParagraphStyle("question", parent=body, fontSize=7, leading=8.5)
    def p(text: str, style=body):
        return Paragraph(escape(text), style)

    def draw_table(rows, widths, top, *, small=False):
        table = Table(rows, colWidths=widths)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#743F2D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7 if small else 8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LINEBELOW", (0, 0), (-1, -1), .25, colors.HexColor("#CBBDB7")),
        ]))
        _, h = table.wrap(width - 64, height)
        if top - h < 35:
            raise ValueError("PDF table exceeds the page; revise layout")
        table.drawOn(canvas, 32, top - h)
        return top - h

    for page in (1, 2, 3):
        canvas.setFillColor(colors.HexColor("#743F2D"))
        canvas.rect(32, height - 69, 42, 37, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 17)
        canvas.drawString(39, height - 55, "KB")
        canvas.setFillColor(colors.HexColor("#743F2D"))
        canvas.setFont("Times-Bold", 17)
        canvas.drawString(85, height - 47, company(data, "V2"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(85, height - 63, "Commercial quotation | " + submission(data, "V2")["received_on"])
        rows = [["Item", "Description and offered rate", "Quantity", "Amount INR"]]
        for ln in data["rfx_lines"]:
            n = ln["line_no"]
            if 1 + (n - 1) // 12 != page:
                continue
            price = field(data, "V2", n, "unit_price")
            total = field(data, "V2", n, "line_total")
            rows.append([str(n), p(price["snippet"]), f"{ln['annual_qty']:,}", total["snippet"]])
        bottom = draw_table(rows, [30, 330, 65, width - 64 - 425], height - 84)
        if page == 2:
            foot = next(r["snippet"] for r in vendor_rows(data, "conditions", "V2")
                        if r["kind"] == "discount_footnote")
            payment = next(r["snippet"] for r in vendor_rows(data, "conditions", "V2")
                           if r["kind"] == "payment_terms")
            para = p(foot + " " + payment, tiny)
            _, h = para.wrap(width - 64, 40)
            para.drawOn(canvas, 32, bottom - h - 6)
            top = bottom - h - 22
            canvas.setFont("Helvetica-Bold", 9)
            canvas.drawString(32, top, "Supplier questionnaire")
            qrows = [["Ref", "Question", "Our response"]]
            for ans in vendor_rows(data, "questionnaire_answers", "V2"):
                qrows.append([str(ans["q_no"]), p(ans["question"], question), p(ans["answer_text"], question)])
            draw_table(qrows, [30, 272, width - 64 - 302], top - 8, small=True)
        if page == 3:
            y = bottom - 30
            for term in vendor_rows(data, "conditions", "V2"):
                if term["kind"] == "discount_footnote":
                    continue
                para = p(term["snippet"])
                _, h = para.wrap(width - 64, 60)
                para.drawOn(canvas, 32, y - h)
                y -= h + 12
            canvas.setFont("Times-Italic", 11)
            canvas.drawString(32, y - 20, "For " + company(data, "V2"))
            canvas.drawString(32, y - 38, "Authorised signatory")
        canvas.showPage()
    canvas.save()


def make_docx(data: dict, out: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.keep_together = True
    doc.styles["Title"].font.name = "Times New Roman"
    doc.styles["Title"].font.size = Pt(22)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    stamp = datetime.fromisoformat(submission(data, "V3")["received_on"])
    doc.core_properties.created = doc.core_properties.modified = stamp
    doc.core_properties.author = company(data, "V3")
    doc.add_paragraph(company(data, "V3") + "\nOffer for corrugated cartons", "Title")  # para1
    doc.add_paragraph("Dear Purchase Team,\nPlease find our size-wise offer below. "
                      + submission(data, "V3")["received_on"])  # para2
    declined = []
    for ln in sorted(data["rfx_lines"], key=lambda x: x["line_no"]):
        n = ln["line_no"]
        row = field(data, "V3", n, "unit_price")
        if row["value"] is None:
            declined.append(ln)
            continue
        if n > 26 and declined:
            chosen = declined[0]
            doc.add_paragraph(f"For the {dims(chosen)} {chosen['ply']}-ply carton, "
                              "we are unable to take this size up at present.")  # para29
        para = doc.add_paragraph()
        para.add_run(dims(ln) + " mm\n").bold = True
        para.add_run(row["snippet"])
        for tier in vendor_rows(data, "tier_rules", "V3"):
            if tier["rfx_line_no"] == n:
                para.add_run(" " + tier["snippet"])
    terms = {r["kind"]: r["snippet"] for r in vendor_rows(data, "conditions", "V3")}
    doc.add_paragraph(terms["validity"])  # para31
    doc.add_paragraph(terms["payment_terms"] + " " + terms["freight"])  # para32
    doc.add_paragraph(terms["moq"])  # para33
    title = doc.add_paragraph("Regarding your supplier questions", "Heading 1")
    title.paragraph_format.page_break_before = False
    for answer in vendor_rows(data, "questionnaire_answers", "V3"):
        para = doc.add_paragraph()
        para.add_run(answer["question"] + " ").italic = True
        para.add_run(answer["answer_text"])
    doc.add_paragraph("Regards,\nCommercial desk\n" + company(data, "V3"))
    path = out / "V3_response.docx"
    doc.save(path)
    stable_zip(path, stamp)


def make_email(data: dict, out: Path) -> None:
    rows = vendor_rows(data, "bid_fields", "V5")
    snippets = {r["snippet"] for r in rows if r["field_name"] == "unit_price" and r["value"] is not None}
    if len(snippets) != 1:
        raise ValueError("V5 truth no longer has a single commercial sentence")
    name = company(data, "V5")
    text = (f"From: Sales Desk | {name}\nTo: Purchase Team\n"
            f"Date: {submission(data, 'V5')['received_on']}\n"
            "Subject: Re: Carton rate enquiry\n\n"
            f"{next(iter(snippets))}\n\nRegards,\nSales & Dispatch\n{name}\n"
            "Please reply on this mail for order confirmation.\n")
    (out / "V5_response.txt").write_text(text, encoding="utf-8", newline="\n")


def validate_outputs(data: dict, out: Path) -> None:
    """Check source values and format traps without changing or repairing truth."""
    from openpyxl import load_workbook
    from docx import Document
    wb = load_workbook(out / "V1_response.xlsx")
    assert wb.sheetnames == ["Quotation", "Terms & Conditions", "Board Specification", "Questionnaire"]
    ws = wb["Quotation"]
    assert all(c.value is None for row in ws.iter_rows(min_row=8, max_row=9) for c in row)
    for ln in data["rfx_lines"]:
        n = ln["line_no"]
        for name in ("unit_price", "line_total", "declared_liner_gsm"):
            r = field(data, "V1", n, name)
            cell = ws[r["anchor"].split("!")[1]]
            assert cell.value == (r["value"] if name == "line_total" else r["snippet"])
        assert ws.cell(n + 9, 8).data_type == "s"
        assert ws.cell(n + 9, 1).value != ln["sku"]
    assert {ws.cell(n + 9, 9).value for n in range(1, 31)} == {"Nos", "nos.", "PCS"}
    paragraphs = Document(out / "V3_response.docx").paragraphs
    text = "\n".join(p.text for p in paragraphs)
    for row in vendor_rows(data, "bid_fields", "V3"):
        if row["field_name"] == "unit_price" and row["value"] is not None:
            index = int(row["anchor"].removeprefix("para")) - 1
            assert row["snippet"] in paragraphs[index].text
    missing = [ln for ln in data["rfx_lines"] if field(data, "V3", ln["line_no"], "unit_price")["value"] is None]
    assert dims(missing[0]) in text
    assert all(dims(ln) not in text for ln in missing[1:])
    assert not Document(out / "V3_response.docx").tables
    for vendor in ("V1", "V2"):
        mismatches = []
        for ln in data["rfx_lines"]:
            price = field(data, vendor, ln["line_no"], "unit_price")
            total = field(data, vendor, ln["line_no"], "line_total")
            expected = round(price["value"] * ln["annual_qty"] /
                             (100 if price["basis"] == "per_100_pieces" else 1), 2)
            if abs(expected - total["value"]) > .01:
                assert total["reason_code"] == "arith_mismatch"
                mismatches.append(ln["line_no"])
        assert len(mismatches) == 1
    for vendor in ("V1", "V2", "V3", "V5"):
        quoted = sum(r["field_name"] == "unit_price" and r["value"] is not None
                     for r in vendor_rows(data, "bid_fields", vendor))
        assert quoted == data["outcomes"]["lines_quoted"][vendor]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed", type=int, default=SEED)
    args = parser.parse_args()
    truth_dir = ROOT / "dataset/truth"
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in truth_dir.iterdir() if p.is_file()}
    data = read_truth(ROOT)
    out = ROOT / "dataset/artifacts"
    out.mkdir(parents=True, exist_ok=True)
    make_excel(data, out, random.Random(args.seed))
    make_pdf(data, out)
    make_docx(data, out)
    make_email(data, out)
    validate_outputs(data, out)
    after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in truth_dir.iterdir() if p.is_file()}
    if before != after:
        raise RuntimeError("Ground truth changed during rendering")
    print("Generated and verified V1_response.xlsx, V2_response.pdf, V3_response.docx, V5_response.txt")
    print("Truth preserved. V1: four slabs; V2: 30 days net; V3: one regret, two silent omissions.")


if __name__ == "__main__":
    main()

