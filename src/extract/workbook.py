"""Deterministic workbook access: the text dump the model reads, and the cell lookup the verifier checks it against.
Both use the same display() so a verbatim snippet is verbatim by construction."""
import re

import openpyxl
from openpyxl.utils import get_column_letter


def split_anchor(anchor):
    sheet, _, ref = anchor.rpartition("!")
    return sheet, ref


def display(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, str):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = cell.number_format
        if f.startswith("#,##0"):
            return f"{v:,.{2 if f.endswith('0.00') else 0}f}"
        return str(v) if isinstance(v, int) else (f"{v:.10f}".rstrip("0").rstrip("."))
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def kind(cell):
    return {"s": "text", "n": "number", "f": "formula", "d": "date", "b": "bool"}.get(cell.data_type, cell.data_type)


class Workbook:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path)

    def cell(self, anchor):
        sheet, ref = split_anchor(anchor)
        if sheet not in self.wb.sheetnames or not re.fullmatch(r"[A-Z]{1,3}[0-9]{1,5}", ref):
            return None
        return self.wb[sheet][ref]

    def text_at(self, anchor):
        c = self.cell(anchor)
        return None if c is None else display(c)

    def dump(self, focus=None, size_pattern=None):
        """One line per non-empty row. With focus={sheet: {row numbers}}, data rows (rows holding a size cell) on that sheet
        are dropped unless listed; headers, terms and every other sheet stay."""
        out = []
        for ws in self.wb:
            merged = ", ".join(str(m) for m in ws.merged_cells.ranges)
            out.append(f"## Sheet '{ws.title}' dimensions {ws.dimensions}" + (f" merged: {merged}" if merged else ""))
            for row in ws.iter_rows():
                cells = [c for c in row if c.value is not None]
                if not cells:
                    continue
                r = cells[0].row
                if focus and ws.title in focus and r not in focus[ws.title] and any(
                        isinstance(c.value, str) and size_pattern.search(c.value) for c in cells):
                    continue
                out.append(f"row {r}: " + " | ".join(f"{c.coordinate} [{kind(c)}] {display(c)!r}" for c in cells))
        return "\n".join(out)

    kind = "xlsx"

    def canonical(self, anchor, sheet=None):
        a = anchor.strip()
        return f"{sheet}!{a}" if sheet and "!" not in a else a

    def verbatim(self, anchor, snippet):
        return (self.text_at(anchor) or "").strip() == (snippet or "").strip()

    def row_of(self, anchor):
        sheet, ref = split_anchor(anchor)
        m = re.fullmatch(r"[A-Z]{1,3}([0-9]{1,5})", ref)
        return (sheet, int(m.group(1))) if m else None

    def is_text_number(self, anchor):
        c = self.cell(anchor)
        return c is not None and c.data_type == "s"

    def candidate_rows(self, sheets):
        from .sources import SIZE
        return {(s, n) for s in sheets for n in self.size_rows(s, SIZE)}

    def size_rows(self, sheet, pattern):
        """Rows on a sheet holding a cell that matches the size pattern: the candidate data rows, found without the model."""
        ws = self.wb[sheet]
        return {c.row for row in ws.iter_rows() for c in row if isinstance(c.value, str) and pattern.search(c.value)}


def col_letter(i):
    return get_column_letter(i)
